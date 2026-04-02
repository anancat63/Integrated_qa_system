from __future__ import annotations

import os
from io import BytesIO
from typing import Iterator, List, Optional

import numpy as np
from PIL import Image
from pptx import Presentation
from tqdm import tqdm

from .ocr import get_ocr
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document


class OCRPPTLoader(BaseLoader):
    SHAPE_TYPE_PICTURE = 13
    SHAPE_TYPE_GROUP = 6

    def __init__(
        self,
        filepath: str,
        *,
        include_position: bool = False,
        include_slide_title: bool = True,
        add_page_break: bool = True,
        ocr_enabled: bool = True,
        ocr_min_image_side: int = 0,
    ) -> None:
        self.filepath = filepath
        self.include_position = include_position
        self.include_slide_title = include_slide_title
        self.add_page_break = add_page_break
        self.ocr_enabled = ocr_enabled
        self.ocr_min_image_side = ocr_min_image_side

    def lazy_load(self) -> Iterator[Document]:
        text = self.ppt2text(self.filepath)
        yield Document(page_content=text, metadata={"source": self.filepath})

    def ppt2text(self, filepath: str) -> str:
        prs = Presentation(filepath)
        ocr = get_ocr() if self.ocr_enabled else None

        out_lines: List[str] = []
        pbar = tqdm(total=len(prs.slides), desc="OCRPPTLoader slide index: 1")

        for slide_number, slide in enumerate(prs.slides, start=1):
            pbar.set_description(f"OCRPPTLoader slide index: {slide_number}")
            pbar.refresh()

            out_lines.append("")
            out_lines.append(f"========== Slide {slide_number} ==========")

            if self.include_slide_title:
                title = self._try_get_slide_title(slide)
                if title:
                    out_lines.append(f"[Slide Title]\n{title}")

            sorted_shapes = sorted(
                slide.shapes,
                key=lambda x: (getattr(x, "top", 0), getattr(x, "left", 0)),
            )
            for shape in sorted_shapes:
                self._extract_shape(shape=shape, out_lines=out_lines, ocr=ocr)

            if self.add_page_break:
                out_lines.append("")

            pbar.update(1)

        text = "\n".join(self._rstrip_empty_lines(out_lines))
        return text

    def _extract_shape(self, shape, out_lines: List[str], ocr) -> None:
        if getattr(shape, "shape_type", None) == self.SHAPE_TYPE_GROUP:
            if self.include_position:
                out_lines.append(self._format_shape_header("[Group]", shape))
            for child in shape.shapes:
                self._extract_shape(child, out_lines, ocr)
            return

        if getattr(shape, "has_text_frame", False):
            txt = (getattr(shape, "text", "") or "").strip()
            if txt:
                out_lines.append(self._format_shape_header("[Text]", shape))
                out_lines.append(txt)
                out_lines.append("")

        if getattr(shape, "has_table", False):
            out_lines.append(self._format_shape_header("[Table]", shape))
            table_lines = self._extract_table_text(shape)
            if table_lines:
                out_lines.extend(table_lines)
            out_lines.append("")

        if getattr(shape, "shape_type", None) == self.SHAPE_TYPE_PICTURE:
            if not self.ocr_enabled or ocr is None:
                return

            try:
                blob = shape.image.blob
            except Exception:
                return

            try:
                image = Image.open(BytesIO(blob))
            except Exception:
                return

            if self.ocr_min_image_side > 0:
                w, h = image.size
                if min(w, h) < self.ocr_min_image_side:
                    return

            result, _ = ocr(np.array(image))
            if result:
                out_lines.append(self._format_shape_header("[Image OCR]", shape))
                ocr_lines = [line[1] for line in result if len(line) > 1 and str(line[1]).strip()]
                if ocr_lines:
                    out_lines.extend(ocr_lines)
                out_lines.append("")

    def _extract_table_text(self, shape) -> List[str]:
        lines: List[str] = []
        try:
            table = shape.table
        except Exception:
            return lines

        for row in table.rows:
            row_cells: List[str] = []
            for cell in row.cells:
                cell_text = (getattr(cell, "text", "") or "").strip()
                row_cells.append(cell_text)
            lines.append("\t".join(row_cells).rstrip())
        return [ln for ln in lines if ln.strip()]

    def _try_get_slide_title(self, slide) -> Optional[str]:
        try:
            if slide.shapes.title is not None:
                title_text = (slide.shapes.title.text or "").strip()
                return title_text if title_text else None
        except Exception:
            pass
        return None

    def _format_shape_header(self, tag: str, shape) -> str:
        if not self.include_position:
            return tag

        top = getattr(shape, "top", None)
        left = getattr(shape, "left", None)
        width = getattr(shape, "width", None)
        height = getattr(shape, "height", None)

        return f"{tag} (top={top}, left={left}, width={width}, height={height})"

    @staticmethod
    def _rstrip_empty_lines(lines: List[str]) -> List[str]:
        i = len(lines) - 1
        while i >= 0 and (lines[i] is None or str(lines[i]).strip() == ""):
            i -= 1
        return lines[: i + 1]


if __name__ == "__main__":
    samples_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "samples"))
    ppt_path = os.path.join(samples_dir, "ocr_01.pptx")

    loader = OCRPPTLoader(
        filepath=ppt_path,
        include_position=False,
        include_slide_title=True,
        add_page_break=True,
        ocr_enabled=True,
        ocr_min_image_side=0,
    )

    docs = loader.load()
    print(type(docs))
    print(docs[0].page_content)

