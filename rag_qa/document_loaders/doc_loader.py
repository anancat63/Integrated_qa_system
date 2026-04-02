import os
from io import BytesIO
from typing import Iterator

import numpy as np
from PIL import Image
from tqdm import tqdm

from .ocr import get_ocr
from docx import Document as Docu1
from docx import ImagePart
from docx.document import Document as Docu2
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document


class OCRDOCLoader(BaseLoader):
    def __init__(self, filepath: str) -> None:
        self.filepath = filepath

    def lazy_load(self) -> Iterator[Document]:
        line = self.doc2text(self.filepath)
        yield Document(page_content=line, metadata={"source": self.filepath})

    def doc2text(self, filepath: str) -> str:
        ocr = get_ocr()
        doc = Docu1(filepath)
        resp = ""

        def iter_block_items(parent):
            if isinstance(parent, Docu2):
                parent_elm = parent.element.body
            elif isinstance(parent, _Cell):
                parent_elm = parent._tc
            else:
                raise ValueError("OCRDOCLoader parse fail")

            for child in parent_elm.iterchildren():
                if isinstance(child, CT_P):
                    yield Paragraph(child, parent)
                elif isinstance(child, CT_Tbl):
                    yield Table(child, parent)

        b_unit = tqdm(total=len(doc.paragraphs) + len(doc.tables), desc="OCRDOCLoader block index: 0")

        for i, block in enumerate(iter_block_items(doc)):
            b_unit.set_description("OCRDOCLoader  block index: {}".format(i))
            b_unit.refresh()

            if isinstance(block, Paragraph):
                resp += block.text.strip() + "\n"
                images = block._element.xpath(".//pic:pic")
                for image in images:
                    for img_id in image.xpath(".//a:blip/@r:embed"):
                        part = doc.part.related_parts[img_id]
                        if isinstance(part, ImagePart):
                            image = Image.open(BytesIO(part._blob))
                            result, _ = ocr(np.array(image))
                            if result:
                                ocr_result = [line[1] for line in result]
                                resp += "\n".join(ocr_result)
            elif isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            resp += paragraph.text.strip() + "\n"

            b_unit.update(1)

        return resp


if __name__ == "__main__":
    samples_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "samples"))
    docx_path = os.path.join(samples_dir, "ocr_02.docx")

    docx_loader = OCRDOCLoader(filepath=docx_path)
    doc = docx_loader.load()
    print(doc)

